"""Decisiones sobre la cola de candidatos tomadas en un CSV o Excel y luego importadas."""

import io
from datetime import UTC, datetime

import pytest
from openpyxl import load_workbook

from ingesta.almacen import ColaCandidatos
from ingesta.candidatos import importar_decisiones
from ingesta.extraccion import Candidato
from ingesta.panel.datos import (
    a_csv,
    a_xlsx,
    candidatos_embedding,
    candidatos_lexicos,
    filas_para_decidir,
    leer_tabla,
    total_candidatos,
)
from ingesta.vocabulario import Vocabulario, guardar_version

from .utiles import VOCAB_PRUEBA

pytestmark = pytest.mark.integracion
NUBE = "Cloud, Infraestructura y Operaciones"


@pytest.fixture
def carpeta(tmp_path):
    guardar_version(VOCAB_PRUEBA, tmp_path)
    return tmp_path


def _encolar(bd, id_aviso, tramo, origen="lexico"):
    c = (
        Candidato(0, f"Uso de {tramo}", tramo, None, None, origen="lexico")
        if origen == "lexico"
        else Candidato(0, f"Uso de {tramo}", tramo, "Alta disponibilidad", 0.85)
    )
    ColaCandidatos(bd).encolar("Bumeran", id_aviso, c, 1, datetime.now(UTC))


def _exportadas(bd, lexicos=True):
    pendientes = candidatos_lexicos(bd, 1) if lexicos else candidatos_embedding(bd)
    return {f["tramo"]: f for f in filas_para_decidir(pendientes, lexicos)}


def _estado(bd, tramo):
    return {c["estado"] for c in bd.candidatos_competencia.find({"tramo": tramo})}


def test_cp44_lote_completo_en_una_sola_version_del_vocabulario(bd, carpeta):
    for id_aviso in ("1", "2"):
        _encolar(bd, id_aviso, "Lambda")
    _encolar(bd, "1", "Efitec")
    _encolar(bd, "1", "Bizagi")
    _encolar(bd, "3", "réplicas", origen="embedding")
    filas = _exportadas(bd)
    filas["Lambda"] |= {"decision": "nueva", "nuevo_termino": "AWS Lambda", "categoria": NUBE}
    filas["Efitec"]["decision"] = "rechazar"
    # Bizagi queda sin decisión: sigue pendiente.
    replica = _exportadas(bd, lexicos=False)["réplicas"]
    assert replica["termino"] == "Alta disponibilidad"  # prellenado con el término cercano
    replica |= {"decision": "variante", "descriptiva": "si"}

    lote = [*filas.values(), replica]
    plan = importar_decisiones(bd, lote, carpeta, aplicar=False)
    assert plan.cuantas("ok") == 3 and _estado(bd, "Lambda") == {"pendiente"}

    hecho = importar_decisiones(bd, lote, carpeta, aplicar=True)
    vocab = Vocabulario.cargar_vigente(carpeta)
    assert hecho.version.name == "vocabulario_v002.json" and vocab.version == 2
    assert vocab.termino("AWS Lambda")["literales"] == ["Lambda"]
    assert "réplicas" in vocab.termino("Alta disponibilidad")["descriptivas"]
    assert _estado(bd, "Lambda") == {"aprobado"}  # los dos avisos
    assert _estado(bd, "Efitec") == {"rechazado"}
    assert _estado(bd, "Bizagi") == {"pendiente"}
    assert bd.lexicos_rechazados.count_documents({"clave": "efitec"}) == 1


def test_cp44_errores_y_omitidas_no_bloquean_las_validas(bd, carpeta):
    for tramo in ("Lambda", "Bizagi", "Efitec", "Signavio", "Dapper"):
        _encolar(bd, "1", tramo)
    filas = _exportadas(bd)
    filas["Lambda"] |= {"decision": "variante", "termino": "No existe"}
    filas["Bizagi"] |= {"decision": "nueva", "categoria": ""}
    filas["Efitec"] |= {"decision": "aprobar"}
    filas["Signavio"] |= {"decision": "nueva", "categoria": "Inventada"}
    filas["Dapper"] |= {"decision": "variante", "termino": "Docker"}
    lote = [*filas.values(), {"id": "123", "decision": "rechazar"}, dict(filas["Dapper"])]
    plan = importar_decisiones(bd, lote, carpeta, aplicar=True)
    detalle = {f.tramo or f.id: (f.estado, f.detalle) for f in plan.filas}
    assert detalle["Lambda"][0] == "error" and "No existe" in detalle["Lambda"][1]
    assert detalle["Bizagi"] == ("error", "una habilidad nueva necesita su categoría")
    assert "decisión desconocida" in detalle["Efitec"][1]
    assert "Inventada" in detalle["Signavio"][1]
    assert detalle["123"] == ("error", "no existe un candidato con ese id")
    assert plan.cuantas("ok") == 1 and plan.cuantas("error") == 6  # Dapper repetido es error
    assert _estado(bd, "Dapper") == {"aprobado"} and _estado(bd, "Lambda") == {"pendiente"}

    otra = importar_decisiones(bd, [filas["Dapper"]], carpeta, aplicar=False)
    assert [(f.estado, f.detalle) for f in otra.filas] == [
        ("omitida", "ya fue revisado (aprobado)")
    ]


def test_cp44_la_misma_habilidad_nueva_dos_veces_se_crea_una_vez(bd, carpeta):
    _encolar(bd, "1", "Lambda")
    _encolar(bd, "1", "lambdas")
    filas = _exportadas(bd)
    for tramo in ("Lambda", "lambdas"):
        filas[tramo] |= {"decision": "nueva", "nuevo_termino": "AWS Lambda", "categoria": NUBE}
    plan = importar_decisiones(bd, list(filas.values()), carpeta, aplicar=True)
    assert plan.cuantas("ok") == 2
    assert "ya existe" in plan.filas[1].detalle
    termino = Vocabulario.cargar_vigente(carpeta).termino("AWS Lambda")
    assert sorted(termino["literales"]) == ["Lambda", "lambdas"]


def test_cp44_ida_y_vuelta_por_excel_con_desplegables(bd, carpeta):
    _encolar(bd, "1", "Lambda")
    filas = list(_exportadas(bd).values())
    contenido = a_xlsx(filas, "candidatos", {"decision": ["variante", "nueva", "rechazar"]})
    libro = load_workbook(io.BytesIO(contenido))
    assert libro["listas"].sheet_state == "hidden"
    assert libro["candidatos"].data_validations.dataValidation[0].formula1 == "=listas!$A$1:$A$3"
    hoja = libro["candidatos"]
    columna = [c.value for c in hoja[1]].index("decision") + 1
    hoja.cell(2, columna, "rechazar")
    salida = io.BytesIO()
    libro.save(salida)
    leidas = leer_tabla("decisiones.xlsx", salida.getvalue())
    assert importar_decisiones(bd, leidas, carpeta, aplicar=True).cuantas("ok") == 1
    assert _estado(bd, "Lambda") == {"rechazado"}


def test_csv_reguardado_por_excel_con_punto_y_coma_y_cp1252(bd, carpeta):
    _encolar(bd, "1", "Lambda")
    fila = next(iter(_exportadas(bd).values()))
    texto = f"id;tramo;decision;categoria\n{fila['id']};Lambda;nueva;{NUBE}\n;;;\n"
    leidas = leer_tabla("decisiones.csv", texto.encode("cp1252"))
    assert len(leidas) == 1
    plan = importar_decisiones(bd, leidas, carpeta, aplicar=False)
    assert [f.estado for f in plan.filas] == ["ok"]
    # El BOM del CSV exportado no ensucia el nombre de la primera columna.
    assert leer_tabla("x.csv", a_csv([fila]))[0]["id"] == fila["id"]


def test_cp44_la_exportacion_no_tiene_el_tope_de_la_tabla(bd):
    for n in range(501):
        _encolar(bd, str(n), f"tramo{n}", origen="embedding")
        _encolar(bd, str(n), f"herramienta{n}")
    assert len(candidatos_embedding(bd)) == 500
    assert len(candidatos_embedding(bd, limite=0)) == 501
    assert len(candidatos_lexicos(bd, 1)) == 500
    assert len(candidatos_lexicos(bd, 1, limite=0)) == 501
    assert total_candidatos(bd, lexicos=False) == total_candidatos(bd, lexicos=True) == 501
