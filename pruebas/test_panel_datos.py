"""Consultas y resaltado del panel."""

import io
from datetime import UTC, datetime

import pytest
from openpyxl import load_workbook

from ingesta.campana import ejecutar_campana
from ingesta.panel.datos import a_csv, a_xlsx, avisos, detalle, resaltar, unidades_resaltadas

from .test_campana import simular_portales


def _c(termino, inicio, fin, via="literal", puntaje=None):
    return {
        "termino": termino,
        "inicio": inicio,
        "fin": fin,
        "via": via,
        "puntaje": puntaje,
        "unidad": 0,
    }


def test_cp46_resaltar_marca_literal_y_semantica_y_escapa_html():
    texto = "Python <avanzado> y backups"
    salida = resaltar(
        texto, [_c("Python", 0, 6), _c("Alta disponibilidad", 20, 27, "semantica", 0.91)]
    )
    assert '<mark class="lit"' in salida and '<mark class="sem"' in salida
    assert "&lt;avanzado&gt;" in salida
    assert "Alta disponibilidad (0.910)" in salida


def test_cp46_resaltar_une_terminos_del_mismo_tramo():
    salida = resaltar("Experiencia con EKS", [_c("AWS", 16, 19), _c("Kubernetes", 16, 19)])
    assert salida.count("<mark") == 1 and "AWS, Kubernetes" in salida


@pytest.mark.integracion
@pytest.mark.respx(assert_all_called=False)
def test_cp46_avisos_y_detalle_desde_las_trazas(respx_mock, bd, dep, params_bow):
    simular_portales(respx_mock, params_bow)
    ejecutar_campana(dep())
    descartes = avisos(bd, destino="descartado")
    assert {f["motivo"] for f in descartes} == {
        "No es de TI",
        "Duplicado",
        "Sin competencias",
        "Rechazado por el esquema",
    }
    con_python = avisos(bd, texto="python")  # Bumeran 1 y su duplicado en Computrabajo
    assert len(con_python) == 2
    assert all("python" in f["título"].lower() for f in con_python)
    d = detalle(bd, "Bumeran", "1")
    assert d["crudo"]["contenido"]["titulo"] == "Desarrollador Python"
    assert d["oferta"]["fuente"]["id_aviso"] == "1" and d["descarte"] is None
    html_unidades = unidades_resaltadas(d["traza"])
    assert any('<mark class="lit"' in h for _, h in html_unidades)


FILAS = [
    {"tramo": "Lambda", "avisos": 3, "ejemplo": "AWS Lambda y SQS", "puntaje": None},
    {"tramo": "Señor", "avisos": 1, "ejemplo": datetime(2026, 10, 4, tzinfo=UTC), "puntaje": 0.8},
]


def test_cp44_csv_con_bom_para_que_excel_respete_las_tildes():
    texto = a_csv(FILAS).decode("utf-8-sig")
    assert texto.splitlines()[0] == "tramo,avisos,ejemplo,puntaje"
    assert "Señor" in texto
    assert a_csv([]) == b""


def test_cp44_xlsx_con_cabecera_y_valores_que_openpyxl_acepta():
    libro = load_workbook(io.BytesIO(a_xlsx(FILAS, "Candidatos")))
    hoja = libro["Candidatos"]
    filas = list(hoja.iter_rows(values_only=True))
    assert filas[0] == ("tramo", "avisos", "ejemplo", "puntaje")
    assert filas[1] == ("Lambda", 3, "AWS Lambda y SQS", None)  # sin caracteres de control
    assert filas[2][2] == "2026-10-04 00:00:00+00:00" and filas[2][3] == 0.8


@pytest.mark.integracion
def test_estado_base_sin_qdrant(bd):
    from qdrant_client import QdrantClient

    from ingesta.panel.datos import estado_base

    caido = QdrantClient(url="http://127.0.0.1:1", timeout=2, check_compatibility=False)
    e = estado_base(bd, caido, "ofertas_pruebas")
    assert e["puntos"] is None
    assert e["error"]
    assert e["ofertas"] == 0


@pytest.mark.integracion
def test_estado_base_con_qdrant(bd, qdrant, params_bc):
    from datetime import UTC, datetime

    from ingesta.conocimiento.sincronizacion import sincronizar
    from ingesta.panel.datos import estado_base

    from .utiles import EmbebedorPalabras, insertar_oferta

    insertar_oferta(bd, "A1")
    sincronizar(bd, qdrant, EmbebedorPalabras(), params_bc.base_conocimiento)
    bd.ejecuciones.insert_one(
        {"modo": "vectorizacion", "estado": "exitosa", "inicio": datetime.now(UTC)}
    )
    e = estado_base(bd, qdrant, "ofertas_pruebas")
    assert (e["ofertas"], e["puntos"], e["error"]) == (1, 1, None)
    assert e["ultima"]["modo"] == "vectorizacion"
