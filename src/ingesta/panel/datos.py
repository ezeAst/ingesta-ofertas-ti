"""Consultas, resaltado y lanzamiento de corridas para el panel (sin dependencias de Streamlit)."""

import csv
import html
import io
import subprocess
from collections import deque
from datetime import datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pymongo.database import Database

from ingesta.config import DIR_LOGS, RAIZ

LIMA = ZoneInfo("America/Lima")
ETIQUETA_MOTIVO = {
    "no_TI": "No es de TI",
    "duplicado": "Duplicado",
    "sin_competencias": "Sin competencias",
    "esquema": "Rechazado por el esquema",
}
_procesos: dict[int, subprocess.Popen] = {}  # corridas lanzadas por este panel


def hora_lima(valor: datetime | None) -> str:
    return valor.astimezone(LIMA).strftime("%Y-%m-%d %H:%M") if valor else ""


def lanzar(argumentos: list[str]) -> int:
    """Inicia `ingesta <argumentos>` como proceso aparte; el panel no se bloquea mientras corre."""
    DIR_LOGS.mkdir(parents=True, exist_ok=True)
    salida = DIR_LOGS / f"panel_{datetime.now(LIMA):%Y-%m-%d_%H%M%S}_{'_'.join(argumentos)}.out"
    with salida.open("w", encoding="utf-8") as archivo:
        proceso = subprocess.Popen(
            ["ingesta", *argumentos],
            cwd=RAIZ,
            stdout=archivo,
            stderr=subprocess.STDOUT,
            start_new_session=True,
        )
    _procesos[proceso.pid] = proceso
    return proceso.pid


def recoger_terminados() -> None:
    """Recoge los procesos terminados para que no queden como zombis."""
    for pid, proceso in list(_procesos.items()):
        if proceso.poll() is not None:
            del _procesos[pid]


def ultimas_lineas(ruta: str | Path | None, n: int = 25) -> str:
    if not ruta or not Path(ruta).exists():
        return ""
    with Path(ruta).open(encoding="utf-8", errors="replace") as archivo:
        return "".join(deque(archivo, maxlen=n))


def avisos(
    bd: Database,
    portal: str | None = None,
    destino: str | None = None,
    motivos: list[str] | None = None,
    texto: str = "",
    limite: int = 1000,
) -> list[dict[str, Any]]:
    """Una fila por aviso procesado, desde su traza. `limite=0`: sin tope (exportación)."""
    filtro: dict[str, Any] = {}
    if portal:
        filtro["portal"] = portal
    if destino:
        filtro["destino"] = destino
    if motivos:
        filtro["motivo"] = {"$in": motivos}
    if texto:
        filtro["titulo"] = {"$regex": _escapar_regex(texto), "$options": "i"}
    proyeccion = {
        "portal": 1,
        "id_aviso": 1,
        "titulo": 1,
        "destino": 1,
        "motivo": 1,
        "fecha": 1,
        "competencias": 1,
        "version_vocabulario": 1,
        "decision_manual": 1,
    }
    filas = []
    for t in bd.trazas.find(filtro, proyeccion).sort("fecha", -1).limit(limite):
        filas.append(
            {
                "portal": t["portal"],
                "id_aviso": t["id_aviso"],
                "título": t.get("titulo", ""),
                "destino": t.get("destino", ""),
                "motivo": ETIQUETA_MOTIVO.get(t.get("motivo"), t.get("motivo") or ""),
                "competencias": len(t.get("competencias", [])),
                "manual": bool(t.get("decision_manual")),
                "vocabulario": t.get("version_vocabulario"),
                "procesado (Lima)": hora_lima(t.get("fecha")),
            }
        )
    return filas


def detalle(bd: Database, portal: str, id_aviso: str) -> dict[str, Any]:
    clave = {"portal": portal, "id_aviso": id_aviso}
    return {
        "crudo": bd.avisos_crudos.find_one(clave, {"_id": 0}),
        "traza": bd.trazas.find_one(clave, {"_id": 0}),
        "oferta": bd.ofertas_laborales.find_one(
            {"fuente.portal": portal, "fuente.id_aviso": id_aviso}, {"_id": 0}
        ),
        "descarte": bd.ofertas_descartadas.find_one(clave, {"_id": 0}),
        "candidatos": list(bd.candidatos_competencia.find(clave, {"_id": 0})),
    }


def unidades_resaltadas(traza: dict[str, Any]) -> list[tuple[str | None, str]]:
    """(sección, HTML) por unidad, con las coincidencias marcadas: literal o semántica."""
    por_unidad: dict[int, list[dict]] = {}
    for c in traza.get("coincidencias", []):
        por_unidad.setdefault(c["unidad"], []).append(c)
    salida = []
    for i, unidad in enumerate(traza.get("unidades", [])):
        salida.append((unidad.get("seccion"), resaltar(unidad["texto"], por_unidad.get(i, []))))
    return salida


def resaltar(texto: str, coincidencias: list[dict]) -> str:
    # Mismo tramo con dos términos (variante colisionada): se muestra una marca con ambos.
    tramos: dict[tuple[int, int], list[dict]] = {}
    for c in coincidencias:
        tramos.setdefault((c["inicio"], c["fin"]), []).append(c)
    partes, cursor = [], 0
    for (inicio, fin), grupo in sorted(tramos.items()):
        if inicio < cursor:  # solapado con el anterior: no se marca dos veces
            continue
        partes.append(html.escape(texto[cursor:inicio]))
        semantica = grupo[0]["via"] == "semantica"
        titulo = ", ".join(
            c["termino"] + (f" ({c['puntaje']:.3f})" if c.get("puntaje") is not None else "")
            for c in grupo
        )
        clase = "sem" if semantica else "lit"
        partes.append(
            f'<mark class="{clase}" title="{html.escape(titulo)}">'
            f"{html.escape(texto[inicio:fin])}</mark>"
            f'<sup class="{clase}">{html.escape(titulo)}</sup>'
        )
        cursor = fin
    partes.append(html.escape(texto[cursor:]))
    return "".join(partes)


def conteos(bd: Database) -> dict[str, int]:
    return {
        "crudos": bd.avisos_crudos.count_documents({}),
        "publicados": bd.ofertas_laborales.count_documents({}),
        "descartados": bd.ofertas_descartadas.count_documents({}),
        "candidatos pendientes": bd.candidatos_competencia.count_documents({"estado": "pendiente"}),
    }


def _pendientes(lexicos: bool, texto: str) -> dict[str, Any]:
    # Los candidatos anteriores a la vía léxica no tienen `origen`: son de embeddings.
    filtro: dict[str, Any] = {
        "estado": "pendiente",
        "origen": "lexico" if lexicos else {"$ne": "lexico"},
    }
    if texto:
        filtro["tramo"] = {"$regex": _escapar_regex(texto), "$options": "i"}
    return filtro


def _grupos_lexicos(min_avisos: int, texto: str) -> list[dict[str, Any]]:
    return [
        {"$match": _pendientes(True, texto)},
        {"$sort": {"creado": 1}},
        {
            "$group": {
                "_id": "$clave",
                "candidato": {"$first": "$_id"},
                "tramo": {"$first": "$tramo"},
                "unidad": {"$first": "$unidad"},
                "portal": {"$first": "$portal"},
                "id_aviso": {"$first": "$id_aviso"},
                "avisos": {"$addToSet": {"$concat": ["$portal", "/", "$id_aviso"]}},
            }
        },
        {"$addFields": {"n_avisos": {"$size": "$avisos"}, "_id": "$candidato"}},
        {"$match": {"n_avisos": {"$gte": min_avisos}}},
    ]


def candidatos_embedding(bd: Database, texto: str = "", limite: int = 500) -> list[dict[str, Any]]:
    """Pendientes de la vía semántica, del puntaje más alto al más bajo.
    `limite=0`: sin tope (exportación)."""
    cursor = bd.candidatos_competencia.find(_pendientes(False, texto)).sort("puntaje", -1)
    return list(cursor.limit(limite))


def candidatos_lexicos(
    bd: Database, min_avisos: int, texto: str = "", limite: int = 500
) -> list[dict[str, Any]]:
    """Pendientes de la vía léxica agrupados por tramo, con la cantidad de avisos que lo mencionan.
    Cada grupo conserva un candidato representativo (`_id`, unidad, aviso) para revisarlo.
    `limite=0`: sin tope (exportación)."""
    etapas = [
        *_grupos_lexicos(min_avisos, texto),
        {"$sort": {"n_avisos": -1, "tramo": 1}},
        *([{"$limit": limite}] if limite else []),
    ]
    return list(bd.candidatos_competencia.aggregate(etapas, allowDiskUse=True))


def total_candidatos(bd: Database, lexicos: bool, min_avisos: int = 1, texto: str = "") -> int:
    """Cuántas filas tendría la exportación: candidatos, o tramos en el caso léxico."""
    if not lexicos:
        return bd.candidatos_competencia.count_documents(_pendientes(False, texto))
    etapas = [*_grupos_lexicos(min_avisos, texto), {"$count": "n"}]
    resultado = list(bd.candidatos_competencia.aggregate(etapas, allowDiskUse=True))
    return resultado[0]["n"] if resultado else 0


def a_csv(filas: list[dict[str, Any]]) -> bytes:
    """CSV con BOM: Excel lo abre en UTF-8 y respeta tildes y eñes."""
    if not filas:
        return b""
    salida = io.StringIO()
    escritor = csv.DictWriter(salida, fieldnames=list(filas[0]))
    escritor.writeheader()
    escritor.writerows(filas)
    return salida.getvalue().encode("utf-8-sig")


def a_xlsx(
    filas: list[dict[str, Any]], hoja: str, listas: dict[str, list[str]] | None = None
) -> bytes:
    """Excel con la cabecera fija. `listas` agrega desplegables por columna; sus valores viven en
    una hoja oculta porque una lista escrita en la fórmula no admite más de 255 caracteres."""
    libro = Workbook()
    ws = libro.active
    ws.title = hoja
    if filas:
        columnas = list(filas[0])
        ws.append(columnas)
        for fila in filas:
            ws.append([_celda(v) for v in fila.values()])
        ws.freeze_panes = "A2"
        if listas:
            opciones = libro.create_sheet("listas")
            opciones.sheet_state = "hidden"
            for n, (columna, valores) in enumerate(listas.items(), start=1):
                letra = get_column_letter(n)
                for fila_lista, valor in enumerate(valores, start=1):
                    opciones.cell(fila_lista, n, _celda(valor))
                validacion = DataValidation(
                    type="list",
                    formula1=f"=listas!${letra}$1:${letra}${max(len(valores), 1)}",
                    allow_blank=True,
                )
                destino = get_column_letter(columnas.index(columna) + 1)
                validacion.add(f"{destino}2:{destino}{len(filas) + 1}")
                ws.add_data_validation(validacion)
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def leer_tabla(nombre: str, contenido: bytes) -> list[dict[str, Any]]:
    """Filas de un CSV o de la primera hoja de un Excel. El CSV puede venir con comas o con punto
    y coma, y en UTF-8 o en la codificación de Excel en Windows (cp1252)."""
    if nombre.lower().endswith(".xlsx"):
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        filas = libro.worksheets[0].iter_rows(values_only=True)
        cabecera = next(filas, None) or ()
        return [
            dict(zip(cabecera, fila, strict=False))
            for fila in filas
            if any(v not in (None, "") for v in fila)
        ]
    try:
        texto = contenido.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = contenido.decode("cp1252")
    if not texto.strip():
        return []
    primera = texto.splitlines()[0]
    separador = ";" if primera.count(";") > primera.count(",") else ","
    return [
        fila
        for fila in csv.DictReader(io.StringIO(texto), delimiter=separador)
        if any((v or "").strip() for v in fila.values() if isinstance(v, str))
    ]


def filas_para_decidir(pendientes: list[dict[str, Any]], lexicos: bool) -> list[dict[str, Any]]:
    """Cola exportable con las columnas que lee `importar_decisiones`, ya prellenadas con lo
    probable: el término cercano (embeddings) o el tramo como habilidad nueva (léxicos)."""
    filas = []
    for c in pendientes:
        if lexicos:
            contexto = {"tramo": c["tramo"], "avisos": c["n_avisos"], "ejemplo": c["unidad"]}
        else:
            contexto = {
                "tramo": c["tramo"],
                "término cercano": c["termino_cercano"],
                "puntaje": c["puntaje"],
                "unidad": c["unidad"],
            }
        filas.append(
            {
                "id": str(c["_id"]),
                **contexto,
                "portal": c["portal"],
                "id_aviso": c["id_aviso"],
                "decision": "",
                "termino": "" if lexicos else c["termino_cercano"],
                "descriptiva": "no",
                "nuevo_termino": c["tramo"] if lexicos else "",
                "categoria": "",
            }
        )
    return filas


def _celda(valor: Any) -> Any:
    """openpyxl rechaza fechas con zona horaria, objetos y caracteres de control."""
    if valor is None or isinstance(valor, bool | int | float):
        return valor
    return ILLEGAL_CHARACTERS_RE.sub("", str(valor))


def _escapar_regex(texto: str) -> str:
    import re

    return re.escape(texto)


def fijar_via_semantica(ruta: Path, activa: bool) -> None:
    """Cambia solo la línea `usar_via_semantica` del TOML; los comentarios se conservan."""
    import re

    texto = ruta.read_text(encoding="utf-8")
    nuevo, n = re.subn(
        r"^usar_via_semantica\s*=\s*(true|false)",
        f"usar_via_semantica = {'true' if activa else 'false'}",
        texto,
        flags=re.MULTILINE,
    )
    if n != 1:
        raise ValueError("no se encontró `usar_via_semantica` en los parámetros")
    ruta.write_text(nuevo, encoding="utf-8", newline="\n")
