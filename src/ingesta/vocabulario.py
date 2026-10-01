"""Vocabulario versionado (Anexo J §7): exportación desde el xlsx de R 1.3 y versiones en JSON."""

import hashlib
import json
import re
from copy import deepcopy
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

# Ramas de la taxonomía de R 1.2: 19 subcategorías técnicas y 7 transversales.
# El campo `tipo` de cada competencia sale de aquí (decisión I14).
_TECNICAS = (
    "Lenguajes de programación",
    "Frameworks y librerías",
    "Bases de datos",
    "Cloud, Infraestructura y Operaciones",
    "Integración y mensajería",
    "Testing y calidad",
    "Seguridad",
    "Control de versiones y CI/CD",
    "Arquitectura y patrones de diseño",
    "Metodologías",
    "Cumplimiento y normativa",
    "Ethical Hacking / Pentesting",
    "Redes y sistemas operativos",
    "Soporte y ticketing",
    "Estilos/UI y diseño visual",
    "Inteligencia Artificial",
    "Analítica, Visualización y Procesamiento de datos",
    "Automatización ofimática",
    "Sistemas Empresariales",
)
_TRANSVERSALES = (
    "Comunicación",
    "Trabajo en equipo",
    "Liderazgo",
    "Proactividad / Iniciativa",
    "Aprendizaje continuo",
    "Pensamiento analítico",
    "Adaptabilidad",
)
RAMAS: dict[str, str] = {c: "técnica" for c in _TECNICAS} | {
    c: "transversal" for c in _TRANSVERSALES
}

HOJA = "Copia"
_RE_ANOTACION = re.compile(r"\s*\([^)]*\)")
_RE_VERSION = re.compile(r"vocabulario_v(\d{3})\.json$")


def _partir(celda: Any) -> list[str]:
    if not celda:
        return []
    vistas: dict[str, None] = {}
    for parte in str(celda).split(","):
        variante = _RE_ANOTACION.sub("", parte).strip()
        if variante:
            vistas.setdefault(variante)
    return list(vistas)


def exportar_desde_xlsx(ruta: Path, hoja: str = HOJA) -> dict[str, Any]:
    libro = load_workbook(ruta, read_only=True, data_only=True)
    filas = libro[hoja].iter_rows(values_only=True)
    cabecera = [str(c).strip() if c else "" for c in next(filas)]
    i_cat = cabecera.index("Categoría")
    i_ter = cabecera.index("Término canónico")
    i_lit = cabecera.index("Variantes y sinónimos")
    i_des = cabecera.index("Variantes descriptivas")
    terminos = []
    for fila in filas:
        if not fila[i_ter]:
            continue
        categoria = str(fila[i_cat]).strip()
        if categoria not in RAMAS:
            raise ValueError(f"categoría fuera de la taxonomía de R 1.2: {categoria!r}")
        terminos.append(
            {
                "nombre": str(fila[i_ter]).strip(),
                "categoria": categoria,
                "literales": _partir(fila[i_lit]),
                "descriptivas": _partir(fila[i_des]),
            }
        )
    usadas = {t["categoria"] for t in terminos}
    return {
        "origen": {
            "archivo": ruta.name,
            "sha256": hashlib.sha256(ruta.read_bytes()).hexdigest(),
            "hoja": hoja,
        },
        "categorias": {c: r for c, r in RAMAS.items() if c in usadas},
        "terminos": terminos,
    }


def _huella(datos: dict[str, Any]) -> str:
    contenido = {"categorias": datos["categorias"], "terminos": datos["terminos"]}
    return hashlib.sha256(json.dumps(contenido, sort_keys=True).encode()).hexdigest()


def _versiones(carpeta: Path) -> list[tuple[int, Path]]:
    return sorted(
        (int(m.group(1)), p)
        for p in carpeta.glob("vocabulario_v*.json")
        if (m := _RE_VERSION.search(p.name))
    )


def guardar_version(datos: dict[str, Any], carpeta: Path) -> Path:
    """Guarda una versión nueva solo si el contenido cambió respecto de la vigente."""
    carpeta.mkdir(parents=True, exist_ok=True)
    versiones = _versiones(carpeta)
    huella = _huella(datos)
    if versiones:
        vigente = json.loads(versiones[-1][1].read_text(encoding="utf-8"))
        if vigente.get("huella") == huella:
            return versiones[-1][1]
    numero = versiones[-1][0] + 1 if versiones else 1
    salida = {
        **datos,
        "version": numero,
        "creado": datetime.now(UTC).isoformat(timespec="seconds"),
        "huella": huella,
    }
    ruta = carpeta / f"vocabulario_v{numero:03d}.json"
    ruta.write_text(json.dumps(salida, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    return ruta


def agregar_variante(
    datos: dict[str, Any], termino: str, variante: str, *, descriptiva: bool
) -> dict[str, Any]:
    nuevo = deepcopy(datos)
    entrada = next((t for t in nuevo["terminos"] if t["nombre"] == termino), None)
    if entrada is None:
        raise ValueError(f"el término {termino!r} no existe en el vocabulario")
    lista = entrada["descriptivas" if descriptiva else "literales"]
    if variante not in lista:
        lista.append(variante)
    return nuevo


def agregar_termino(datos: dict[str, Any], nombre: str, categoria: str) -> dict[str, Any]:
    """Añade un término canónico en una categoría existente: de R 1.2 o creada desde el panel."""
    rama = datos["categorias"].get(categoria) or RAMAS.get(categoria)
    if rama is None:
        raise ValueError(f"la categoría {categoria!r} no existe en el vocabulario")
    nuevo = deepcopy(datos)
    if any(t["nombre"] == nombre for t in nuevo["terminos"]):
        raise ValueError(f"el término {nombre!r} ya existe")
    nuevo["terminos"].append(
        {"nombre": nombre, "categoria": categoria, "literales": [], "descriptivas": []}
    )
    nuevo["categorias"].setdefault(categoria, rama)
    return nuevo


def agregar_categoria(datos: dict[str, Any], nombre: str, rama: str) -> dict[str, Any]:
    """Categoría nueva fuera de la taxonomía de R 1.2 (se documenta al actualizar R 1.2)."""
    if rama not in ("técnica", "transversal"):
        raise ValueError(f"la rama debe ser «técnica» o «transversal», no {rama!r}")
    nombre = nombre.strip()
    if not nombre:
        raise ValueError("la categoría necesita un nombre")
    if nombre in datos["categorias"] or nombre in RAMAS:
        raise ValueError(f"la categoría {nombre!r} ya existe")
    nuevo = deepcopy(datos)
    nuevo["categorias"][nombre] = rama
    return nuevo


class Vocabulario:
    def __init__(self, datos: dict[str, Any]):
        self.datos = datos
        self.version: int = datos.get("version", 0)
        self._por_nombre = {t["nombre"]: t for t in datos["terminos"]}

    @classmethod
    def cargar_vigente(cls, carpeta: Path) -> "Vocabulario":
        versiones = _versiones(carpeta)
        if not versiones:
            raise FileNotFoundError(f"no hay vocabulario exportado en {carpeta}")
        return cls(json.loads(versiones[-1][1].read_text(encoding="utf-8")))

    @property
    def terminos(self) -> list[dict[str, Any]]:
        return self.datos["terminos"]

    def termino(self, nombre: str) -> dict[str, Any]:
        return self._por_nombre[nombre]

    def rama(self, categoria: str) -> str:
        return self.datos["categorias"][categoria]
